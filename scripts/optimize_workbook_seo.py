#!/usr/bin/env python3

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


BRAND = "CAMARI JAPAN"
TITLE_LIMIT = 65
DESCRIPTION_LIMIT = 165

TEXT_REPLACEMENTS = {
    "automtive": "automotive",
    "Automtive": "Automotive",
}


@dataclass
class Change:
    workbook: Path
    sheet: str
    row: int
    key: str
    column: str
    before: str
    after: str


def clean_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\s+", " ", text).strip()

    for before, after in TEXT_REPLACEMENTS.items():
        text = text.replace(before, after)

    text = text.replace(" .", ".").replace(" ,", ",")
    text = text.replace(". for ", " for ")
    text = re.sub(r"\s+\|", " |", text)
    return text


def title_case_color(value: str) -> str:
    if not value:
        return ""

    return " / ".join(part.strip().title() for part in value.split("/") if part.strip())


def truncate_sentence(value: str, limit: int) -> str:
    value = clean_text(value)

    if len(value) <= limit:
        return value

    trimmed = value[: limit + 1]
    for marker in [". ", "; ", ", "]:
        cut = trimmed.rfind(marker)
        if cut >= 80:
            sentence = trimmed[:cut].rstrip(" ,.;")
            return f"{sentence}."

    return trimmed[: limit - 1].rstrip(" ,.;") + "."


def compact_description(value: str, limit: int = DESCRIPTION_LIMIT) -> str:
    value = clean_text(value)

    if len(value) <= limit:
        return value

    candidates = [
        re.sub(r"\s+Available in multiple colours through Camari\.\s*", " ", value, flags=re.IGNORECASE),
        re.sub(r"\s+View specs and downloads\.\s*", " ", value, flags=re.IGNORECASE),
        re.sub(r"\s+More colors available\.\s*", " ", value, flags=re.IGNORECASE),
    ]

    for candidate in candidates:
        candidate = clean_text(candidate)
        if len(candidate) <= limit:
            return candidate

    compressed = clean_text(value)
    compressed = compressed.replace(" with fire retardant option", "")
    compressed = compressed.replace(" and request sample information", "")
    if len(compressed) <= limit:
        return compressed

    return truncate_sentence(compressed, limit)


def sentence_case(value: str) -> str:
    value = clean_text(value)
    if not value:
        return value
    return value[0].upper() + value[1:]


def compact_title(parts: Iterable[str]) -> str:
    cleaned = [clean_text(part) for part in parts if clean_text(part)]
    title = " ".join(cleaned)
    suffix = f" | {BRAND}"

    if not title.endswith(suffix):
        title = f"{title}{suffix}"

    if len(title) <= TITLE_LIMIT:
        return title

    available = TITLE_LIMIT - len(suffix)
    base = " ".join(cleaned)
    base = base[:available].rstrip(" -/,")
    return f"{base}{suffix}"


def build_product_type_description(name: str, summary: str) -> str:
    summary = clean_text(summary)

    if summary:
        return truncate_sentence(summary, DESCRIPTION_LIMIT)

    return truncate_sentence(f"{name} material collection for premium automotive and interior applications. View colours, specifications, downloads, and project uses.", DESCRIPTION_LIMIT)


def material_phrase(material_slug: str, product_type_name: str) -> str:
    name = product_type_name.lower()

    if material_slug == "alcantara":
        return "Italian microfibre surface for automotive interiors and trim"
    if material_slug == "leather" or "leather" in name or "nappa" in name:
        return "premium leather surface for automotive and interior applications"
    if material_slug == "fabric" or "fabric" in name:
        return "automotive upholstery fabric for restoration and interior trim"

    return "premium surface material for automotive and interior applications"


def build_sku_summary(material_slug: str, product_type_name: str, code: str, color: str) -> str:
    color_text = title_case_color(color)
    phrase = sentence_case(material_phrase(material_slug, product_type_name))
    color_clause = f" in {color_text}" if color_text else ""
    return truncate_sentence(f"{product_type_name}, code {code}{color_clause}. {phrase}.", DESCRIPTION_LIMIT)


def build_sku_description(material_slug: str, product_type_name: str, code: str, color: str) -> str:
    color_text = title_case_color(color)
    phrase = sentence_case(material_phrase(material_slug, product_type_name))
    color_clause = f" in {color_text}" if color_text else ""
    return truncate_sentence(f"{product_type_name} {code}{color_clause}. {phrase}. View specifications, colours, and request sample information.", DESCRIPTION_LIMIT)


def build_sku_summary_ja(product_type_name: str, code: str, color: str) -> str:
    color_text = title_case_color(color)
    color_clause = f"、カラー {color_text}" if color_text else ""
    return truncate_sentence(f"{product_type_name}、品番 {code}{color_clause}。レストアおよび内装トリム向けの自動車用マテリアルです。", DESCRIPTION_LIMIT)


def build_sku_description_ja(product_type_name: str, code: str, color: str) -> str:
    color_text = title_case_color(color)
    color_clause = f"、カラー {color_text}" if color_text else ""
    return truncate_sentence(f"{product_type_name} {code}{color_clause}。仕様、カラー、サンプル依頼情報をご確認いただけます。", DESCRIPTION_LIMIT)


def headers_for(sheet) -> dict[str, int]:
    return {str(cell.value).strip(): index for index, cell in enumerate(sheet[1], start=1) if cell.value}


def value_at(sheet, row: int, headers: dict[str, int], column: str) -> str:
    index = headers.get(column)
    if not index:
        return ""
    return clean_text(sheet.cell(row=row, column=index).value)


def maybe_change(
    changes: list[Change],
    workbook: Path,
    sheet_name: str,
    row: int,
    key: str,
    column: str,
    before: str,
    after: str,
) -> None:
    before = "" if before is None else str(before).strip()
    after = clean_text(after)

    if before != after:
        changes.append(Change(workbook, sheet_name, row, key, column, before, after))


def collect_changes(workbook_path: Path) -> list[Change]:
    workbook = load_workbook(workbook_path)
    changes: list[Change] = []

    if "product_types" not in workbook.sheetnames or "skus" not in workbook.sheetnames:
        return changes

    product_type_sheet = workbook["product_types"]
    product_type_headers = headers_for(product_type_sheet)
    product_types: dict[str, dict[str, str]] = {}

    for row in range(2, product_type_sheet.max_row + 1):
        slug = value_at(product_type_sheet, row, product_type_headers, "product_type_slug")
        name_en = value_at(product_type_sheet, row, product_type_headers, "name_en")

        if not slug or not name_en:
            continue

        name_ja = value_at(product_type_sheet, row, product_type_headers, "name_ja")
        summary_en = value_at(product_type_sheet, row, product_type_headers, "summary_en")
        material_slug = value_at(product_type_sheet, row, product_type_headers, "material_slug")
        seo_title_en = value_at(product_type_sheet, row, product_type_headers, "seo_title_en")
        seo_description_en = value_at(product_type_sheet, row, product_type_headers, "seo_description_en")
        seo_title_ja = value_at(product_type_sheet, row, product_type_headers, "seo_title_ja")
        seo_description_ja = value_at(product_type_sheet, row, product_type_headers, "seo_description_ja")

        product_types[slug] = {
            "name_en": name_en,
            "name_ja": name_ja,
            "material_slug": material_slug,
            "summary_en": summary_en,
        }

        suggested_title = clean_text(seo_title_en) if seo_title_en else compact_title([name_en])
        suggested_description = compact_description(seo_description_en) if seo_description_en else build_product_type_description(name_en, summary_en)

        if not seo_title_en or clean_text(seo_title_en) != seo_title_en:
            maybe_change(changes, workbook_path, "product_types", row, slug, "seo_title_en", seo_title_en, suggested_title)
        if not seo_description_en or len(seo_description_en) > DESCRIPTION_LIMIT or clean_text(seo_description_en) != seo_description_en:
            maybe_change(changes, workbook_path, "product_types", row, slug, "seo_description_en", seo_description_en, suggested_description)
        if seo_title_ja and clean_text(seo_title_ja) != seo_title_ja:
            maybe_change(changes, workbook_path, "product_types", row, slug, "seo_title_ja", seo_title_ja, seo_title_ja)
        if seo_description_ja and clean_text(seo_description_ja) != seo_description_ja:
            maybe_change(changes, workbook_path, "product_types", row, slug, "seo_description_ja", seo_description_ja, seo_description_ja)

    sku_sheet = workbook["skus"]
    sku_headers = headers_for(sku_sheet)

    for row in range(2, sku_sheet.max_row + 1):
        sku_slug = value_at(sku_sheet, row, sku_headers, "sku_slug")

        if not sku_slug:
            continue

        product_type_slug = value_at(sku_sheet, row, sku_headers, "product_type_slug")
        product_type = product_types.get(product_type_slug, {})
        product_type_name = product_type.get("name_en", product_type_slug.replace("-", " ").title())
        material_slug = value_at(sku_sheet, row, sku_headers, "material_slug") or product_type.get("material_slug", "")
        code = value_at(sku_sheet, row, sku_headers, "code")
        color_en = value_at(sku_sheet, row, sku_headers, "color_name_en")
        image = value_at(sku_sheet, row, sku_headers, "image")
        seo_image = value_at(sku_sheet, row, sku_headers, "seo_image")
        summary_en = value_at(sku_sheet, row, sku_headers, "summary_en")
        summary_ja = value_at(sku_sheet, row, sku_headers, "summary_ja")
        seo_title_en = value_at(sku_sheet, row, sku_headers, "seo_title_en")
        seo_title_ja = value_at(sku_sheet, row, sku_headers, "seo_title_ja")
        seo_description_en = value_at(sku_sheet, row, sku_headers, "seo_description_en")
        seo_description_ja = value_at(sku_sheet, row, sku_headers, "seo_description_ja")

        suggested_summary = build_sku_summary(material_slug, product_type_name, code, color_en)
        suggested_summary_ja = build_sku_summary_ja(product_type.get("name_ja") or product_type_name, code, color_en)
        suggested_title = compact_title([product_type_name, code, title_case_color(color_en)])
        suggested_title_ja = compact_title([product_type.get("name_ja") or product_type_name, code, title_case_color(color_en)])
        suggested_description = build_sku_description(material_slug, product_type_name, code, color_en)
        suggested_description_ja = build_sku_description_ja(product_type.get("name_ja") or product_type_name, code, color_en)

        if not summary_en or clean_text(summary_en) != summary_en:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "summary_en", summary_en, suggested_summary)
        if not summary_ja or clean_text(summary_ja) != summary_ja:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "summary_ja", summary_ja, suggested_summary_ja)
        if not seo_title_en or clean_text(seo_title_en) != seo_title_en:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "seo_title_en", seo_title_en, suggested_title)
        if not seo_title_ja or clean_text(seo_title_ja) != seo_title_ja:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "seo_title_ja", seo_title_ja, suggested_title_ja)
        if not seo_description_en or len(seo_description_en) > DESCRIPTION_LIMIT or clean_text(seo_description_en) != seo_description_en:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "seo_description_en", seo_description_en, suggested_description)
        if not seo_description_ja or len(seo_description_ja) > DESCRIPTION_LIMIT or clean_text(seo_description_ja) != seo_description_ja:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "seo_description_ja", seo_description_ja, suggested_description_ja)
        if not seo_image and image:
            maybe_change(changes, workbook_path, "skus", row, sku_slug, "seo_image", seo_image, image)

    return changes


def apply_changes(workbook_path: Path, changes: list[Change]) -> None:
    workbook = load_workbook(workbook_path)

    for change in changes:
        sheet = workbook[change.sheet]
        headers = headers_for(sheet)
        column_index = headers.get(change.column)

        if column_index:
            sheet.cell(row=change.row, column=column_index).value = change.after

    workbook.save(workbook_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audit and optimize product SEO fields in catalog import workbooks.")
    parser.add_argument("workbooks", nargs="+", type=Path, help="Workbook paths to audit or update.")
    parser.add_argument("--write", action="store_true", help="Write suggested SEO changes back to the workbook files.")
    parser.add_argument("--limit", type=int, default=12, help="Maximum detailed changes to print per workbook.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    all_changes: list[Change] = []

    for workbook_path in args.workbooks:
        changes = collect_changes(workbook_path)
        all_changes.extend(changes)
        print(f"\n{workbook_path}")
        print(f"  suggested changes: {len(changes)}")

        for change in changes[: args.limit]:
            before = change.before or "(blank)"
            print(f"  - {change.sheet} row {change.row} {change.column} [{change.key}]")
            print(f"    before: {before}")
            print(f"    after:  {change.after}")

        if len(changes) > args.limit:
            print(f"  ... {len(changes) - args.limit} more")

        if args.write and changes:
            apply_changes(workbook_path, changes)
            print("  written")

    mode = "written" if args.write else "dry-run"
    print(f"\nTotal suggested changes ({mode}): {len(all_changes)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
